"""Выгрузка в Excel: сводка по прогонам или сырые сообщения одного прогона.

    python analysis/to_excel.py catalog            → analysis/out/catalog.xlsx
    python analysis/to_excel.py team               → analysis/out/записи_для_команды.xlsx (упрощённая)
    python analysis/to_excel.py 30618_01f73500     → analysis/out/run_30618_01f73500.xlsx
"""
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bagio import DATA_DIR, ROOT, read_bag
from catalog import KMH, OUT_DIR
from map import nearest_terminal, track

FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)
GREY = Font(name='Arial', size=10, color='999999')
HEAD_FILL = PatternFill('solid', fgColor='DDE4EE')
BAD_FILL = PatternFill('solid', fgColor='F8D7D3')

# (столбец catalog.csv, заголовок, формат, пояснение)
COLUMNS = [
    ('bag', 'Прогон', '@', 'Имя папки в dataset/data: <трамвай>_<хеш>.'),
    ('vehicle', 'Трамвай', '0', 'Номер вагона: 30618 или 30639.'),
    ('notes', 'Замечания', '@', 'Краткий список проблем прогона.'),
    ('duplicate_of', 'Копия прогона', '@', 'Если заполнено — прогон байт в байт совпадает с указанным, использовать его отдельно не надо.'),
    ('duration_s', 'Длительность, с', '0', 'От первого до последнего сообщения (по header.stamp).'),
    ('ref_dist_m', 'Путь по GNSS, м', '0', 'Интеграл скорости GNSS по времени. Пусто — GNSS в прогоне нет.'),
    ('front_dist_m', 'Путь по передней тележке, м', '0', 'Интеграл скорости передней тележки (уже переведённой из км/ч в м/с).'),
    ('ref_max_speed', 'Макс. скорость GNSS, м/с', '0.0', 'Максимум скорости по GNSS.'),
    ('stops', 'Остановок', '0', 'Сколько раз скорость GNSS была < 0,2 м/с дольше 3 с.'),
    ('ref_speed_first5s', 'Скорость в первые 5 с, м/с', '0.00', 'Средняя скорость GNSS в первые 5 секунд: ≈ 0 значит, что трамвай стоит.'),
    ('ref_src', 'Эталон', '@', 'Какой приёмник GNSS взят как эталон скорости: rover (передняя антенна), если есть, иначе master.'),
    ('master_rtk_share', 'Точное решение master', '0%', 'Доля сообщений master с status = 2 (высокоточное решение).'),
    ('rover_rtk_share', 'Точное решение rover', '0%', 'То же для rover.'),
    ('master_fix_max_gap_s', 'Макс. пропуск master, с', '0.0', 'Самая длинная пауза между сообщениями master/fix.'),
    ('front_err', 'Колёса врут, перед.', '+0.00%;-0.00%;0.00%', 'На сколько процентов передняя тележка показывает больше (+) или меньше (−) реальной скорости по GNSS. Медиана по моментам, когда трамвай едет быстрее 3 м/с и нет проскальзывания. Причина — колесо стёрлось, и его реальная окружность не равна той, что заложена в датчик.'),
    ('rear_err', 'Колёса врут, зад.', '+0.00%;-0.00%;0.00%', 'То же для задней тележки.'),
    ('front_slip_share', 'Расхождение с GNSS, перед.', '0.00%', 'Доля времени движения, когда передняя тележка отличается от GNSS больше чем на 1 м/с (проскальзывание, юз, сбой).'),
    ('rear_slip_share', 'Расхождение с GNSS, зад.', '0.00%', 'То же для задней.'),
    ('front_rear_mismatch_share', 'Перед ≠ зад', '0.00%', 'Доля времени движения, когда тележки отличаются друг от друга больше чем на 0,5 м/с.'),
    ('front_max_gap_s', 'Макс. пропуск передней, с', '0.0', 'Самая длинная пауза между сообщениями передней тележки (норма ≈ 0,1 с).'),
    ('rear_max_gap_s', 'Макс. пропуск задней, с', '0.0', 'То же для задней.'),
    ('cmd_max_gap_s', 'Макс. пропуск ручки, с', '0.00', 'То же для ручки (норма ≈ 0,05 с).'),
    ('front_late_msgs', 'Опоздавших, перед.', '0', 'Сколько сообщений пришло в запись позже, чем более новые измерения (время измерения меньше, чем у уже пришедшего сообщения).'),
    ('rear_late_msgs', 'Опоздавших, зад.', '0', 'То же для задней.'),
    ('cmd_late_msgs', 'Опоздавших, ручка', '0', 'То же для ручки.'),
    ('front_late_max_s', 'Макс. опоздание, перед., с', '0.00', 'На сколько секунд опоздало самое позднее сообщение.'),
    ('front_spikes', 'Скачков, перед.', '0', 'Скачков скорости > 1 м/с между соседними сообщениями (в порядке времени, не через пропуск).'),
    ('rear_spikes', 'Скачков, зад.', '0', 'То же для задней.'),
    ('front_hz', 'Частота передней, Гц', '0.0', 'Сколько сообщений в секунду (по медиане интервала).'),
    ('cmd_hz', 'Частота ручки, Гц', '0.0', 'То же для ручки.'),
    ('front_lag_ms', 'Задержка записи тележек, мс', '0', 'Медиана (время записи в bag − header.stamp).'),
    ('cmd_traction_share', 'Время в тяге', '0%', 'Доля сообщений ручки с позицией > 0.'),
    ('cmd_brake_share', 'Время в торможении', '0%', 'Доля сообщений ручки с позицией < 0.'),
    ('cmd_changes', 'Переключений ручки', '0', 'Сколько раз позиция ручки менялась.'),
    ('antenna_baseline_m', 'Между антеннами, м', '0.00', 'Медианное расстояние master ↔ rover.'),
]


def notes(r):
    out = []
    if isinstance(r.duplicate_of, str) and r.duplicate_of:
        out.append(f'копия {r.duplicate_of}')
    if r.short:
        out.append('короткий')
    if pd.isna(r.ref_dist_m):
        out.append('нет GNSS')
    elif r.n_master_fix == 0:
        out.append('нет master')
    for k, name in (('front', 'передний'), ('rear', 'задний')):
        if r[f'{k}_max_gap_s'] > 5:
            out.append(f'{name} датчик молчал {r[f"{k}_max_gap_s"]:.0f} с')
        if r[f'{k}_slip_share'] > 0.002:
            out.append(f'{name}: расхождение с GNSS {r[f"{k}_slip_share"]:.1%} времени')
    if max(r.front_late_max_s, r.rear_late_max_s, r.cmd_late_max_s) > 0.5:
        out.append('опоздавшие сообщения до '
                   f'{max(r.front_late_max_s, r.rear_late_max_s, r.cmd_late_max_s):.1f} с')
    return '; '.join(out)


def style_header(ws, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 42


def catalog_xlsx():
    cat = pd.read_csv(OUT_DIR / 'catalog.csv', keep_default_na=False, na_values=[''])
    cat['notes'] = cat.apply(notes, axis=1)
    cat['front_err'], cat['rear_err'] = cat.front_scale - 1, cat.rear_scale - 1
    wb = Workbook()
    ws = wb.active
    ws.title = 'Прогоны'
    ws.append([h for _, h, _, _ in COLUMNS])
    for _, r in cat.iterrows():
        ws.append([None if pd.isna(r[c]) else (r[c].item() if hasattr(r[c], 'item') else r[c])
                   for c, _, _, _ in COLUMNS])
    bad = {'front_max_gap_s': 5, 'rear_max_gap_s': 5, 'front_slip_share': 0.002,
           'rear_slip_share': 0.002, 'front_late_max_s': 0.5}
    for row in ws.iter_rows(min_row=2):
        r = cat.iloc[row[0].row - 2]
        skip = bool(r.short) or bool(isinstance(r.duplicate_of, str) and r.duplicate_of)
        for cell, (col, _, fmt, _) in zip(row, COLUMNS):
            cell.font = GREY if skip else FONT
            cell.number_format = fmt
            if col in bad and cell.value is not None and cell.value > bad[col]:
                cell.fill = BAD_FILL
    style_header(ws, [16, 9, 48, 16] + [12] * (len(COLUMNS) - 4))
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions

    desc = wb.create_sheet('Описание столбцов')
    desc.append(['Столбец', 'Что значит'])
    for _, h, _, text in COLUMNS:
        desc.append([h, text])
    desc.append([])
    desc.append(['Серым', 'Копии других прогонов и прогоны короче 60 с — для анализа их не берём.'])
    desc.append(['Красная заливка', 'Датчик молчал > 5 с; расхождение с GNSS > 0,2 % времени; опоздания > 0,5 с.'])
    desc.append(['Единицы', 'Скорость тележек в bag записана в км/ч; здесь везде переведена в м/с (÷ 3,6).'])
    for row in desc.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            c.alignment = Alignment(wrap_text=True, vertical='top')
    style_header(desc, [30, 100])
    desc.row_dimensions[1].height = 18

    path = OUT_DIR / 'catalog.xlsx'
    wb.save(path)
    return path


def compare_table(d, t0):
    """Все сигналы на одной временной шкале: строка на каждое сообщение передней тележки."""
    def rel(df):
        return df.assign(t=(df.t_hdr - t0) / 1e9).sort_values('t')

    out = rel(d['front'])[['t']].copy()
    out['передняя, м/с'] = rel(d['front']).velocity.values / KMH
    for name, col, src in (('rear', 'задняя, м/с', 'velocity'), ('rover_vel', 'GNSS rover, м/с', None),
                           ('master_vel', 'GNSS master, м/с', None), ('cmd', 'ручка', 'position')):
        if not len(d[name]):
            continue
        df = rel(d[name])
        if name == 'rear':
            df['value'] = df.velocity / KMH
        elif src is None:
            df['value'] = np.hypot(df.vx, df.vy)
        else:
            df['value'] = df[src]
        # ручка — последнее значение до момента; остальное — ближайшее по времени в пределах 0,06 с
        kw = {'direction': 'backward'} if name == 'cmd' else {'direction': 'nearest', 'tolerance': 0.06}
        out = pd.merge_asof(out, df[['t', 'value']].rename(columns={'value': col}), on='t', **kw)
    gnss = 'GNSS rover, м/с' if 'GNSS rover, м/с' in out else 'GNSS master, м/с'
    if gnss in out:
        out['передняя − GNSS'] = out['передняя, м/с'] - out[gnss]
        if 'задняя, м/с' in out:
            out['задняя − GNSS'] = out['задняя, м/с'] - out[gnss]
    return out.rename(columns={'t': 't, с от начала'})


def run_xlsx(bag):
    d = read_bag(DATA_DIR / bag)
    t0 = min(df.t_hdr.min() for df in d.values() if len(df))
    path = OUT_DIR / f'run_{bag}.xlsx'
    try:
        path.open('ab').close()
    except PermissionError:  # файл открыт в Excel — пишем рядом
        path = path.with_name(f'run_{bag}_new.xlsx')
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        cmp = compare_table(d, t0)
        cmp.to_excel(xw, sheet_name='Сравнение', index=False)
        ws = xw.sheets['Сравнение']
        style_header(ws, [12] * len(cmp.columns))
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        diff_cols = [i for i, c in enumerate(cmp.columns, 1) if c.endswith('− GNSS')]
        for row in ws.iter_rows(min_row=2):
            for i in diff_cols:
                v = row[i - 1].value
                if isinstance(v, (int, float)) and abs(v) > 1:  # пустые ячейки pandas пишет как ''
                    row[i - 1].fill = BAD_FILL
            for c in row:
                c.number_format = '0.000' if c.column == 1 else '0.00'

        for name, df in d.items():
            if not len(df):
                continue
            # строки — в порядке записи в bag, то есть в том порядке, в каком их получит нода
            out = pd.DataFrame({
                'время измерения (header.stamp), нс': df.t_hdr.astype(str),
                'время записи в bag, нс': df.t_bag.astype(str),
                'измерение, с от начала': (df.t_hdr - t0) / 1e9,
                'запись, с от начала': (df.t_bag - t0) / 1e9,
                'записано позже измерения на, с': (df.t_bag - df.t_hdr) / 1e9,
            })
            if name in ('front', 'rear'):
                out['velocity как в bag (км/ч)'] = df.velocity
                out['скорость, м/с'] = df.velocity / KMH
            elif name == 'cmd':
                out['позиция ручки'] = df.position
            elif name.endswith('_vel'):
                out['vx (на восток), м/с'] = df.vx
                out['vy (на север), м/с'] = df.vy
                out['скорость, м/с'] = np.hypot(df.vx, df.vy)
            else:
                out['широта'] = df.lat
                out['долгота'] = df.lon
                out['высота, м'] = df.alt
                out['status'] = df.status
            out.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            style_header(ws, [22, 22] + [14] * (len(out.columns) - 2))
            ws.freeze_panes = 'C2'
    return path


TERMINAL_NAMES = {'A': 'Щукинская', 'B': 'Таллинская'}
PURPOSE_ORDER = ['точность', 'точность + устойчивость', 'настройка модели', 'отладка', 'не брать']


def purpose(r):
    if isinstance(r.duplicate_of, str) and r.duplicate_of:
        return 'не брать'
    if r.short:
        return 'отладка'
    if pd.isna(r.ref_dist_m):
        return 'настройка модели'
    problems = (max(r.front_max_gap_s, r.rear_max_gap_s) > 5
                or max(r.front_slip_share, r.rear_slip_share) > 0.002
                or np.nanmax([r.front_late_max_s, r.rear_late_max_s, r.cmd_late_max_s]) > 0.5)
    return 'точность + устойчивость' if problems else 'точность'


def direction(r):
    if pd.isna(r.ref_dist_m) or r.short:
        return ''
    fix = track(r.bag)[0]
    a = nearest_terminal(fix.lat.iloc[0], fix.lon.iloc[0])
    b = nearest_terminal(fix.lat.iloc[-1], fix.lon.iloc[-1])
    return f'{TERMINAL_NAMES[a]} → {TERMINAL_NAMES[b]}'


def purpose_and_notes(cat):
    """Для каждой строки каталога: для чего брать запись и её проблемы простым текстом."""
    p = cat.apply(purpose, axis=1)
    n = (cat.apply(notes, axis=1)
         .str.replace(r'копия [^;]+(; )?|короткий(; )?|нет GNSS(; )?|нет master(; )?', '', regex=True)
         .str.replace(r'(\d)\.(\d)', lambda m: m.group(1) + ',' + m.group(2), regex=True).str.rstrip('; '))
    return p, n.where(p != 'не брать', '')


def team_xlsx():
    """Упрощённая сводка для команды: какую запись для чего брать."""
    cat = pd.read_csv(OUT_DIR / 'catalog.csv', keep_default_na=False, na_values=[''])
    cat['purpose'], cat['notes'] = purpose_and_notes(cat)
    cat['direction'] = cat.apply(direction, axis=1)
    cat['wheel_err'] = (cat.front_scale + cat.rear_scale) / 2 - 1
    cat['order'] = cat.purpose.map(PURPOSE_ORDER.index)
    cat = cat.sort_values(['order', 'bag'])
    cols = [('bag', 'Запись', '@', 16), ('purpose', 'Для чего брать', '@', 24),
            ('vehicle', 'Трамвай', '0', 9), ('direction', 'Направление', '@', 22),
            ('duration_s', 'Длительность, мин', '0.0', 13), ('ref_dist_m', 'Путь по GNSS, км', '0.00', 12),
            ('wheel_err', 'Колёса врут', '+0.0%;-0.0%;0.0%', 11), ('notes', 'Проблемы в записи', '@', 60),
            ('duplicate_of', 'Копия записи', '@', 16)]
    wb = Workbook()
    info = wb.active
    info.title = 'Кратко'
    long_ = cat[~cat.short & (cat.purpose != 'не брать')]
    lines = [
        ('Всего записей', len(cat)),
        ('из них копии других (не брать)', int((cat.purpose == 'не брать').sum())),
        ('короткие, меньше минуты (только для отладки)', int((cat.purpose == 'отладка').sum())),
        ('нормальные длинные записи', len(long_)),
        ('   с GNSS — на них проверяем точность', int(long_.ref_dist_m.notna().sum())),
        ('   из них с проблемами — проверяем устойчивость', int((cat.purpose == 'точность + устойчивость').sum())),
        ('   без GNSS — только для настройки модели', int((cat.purpose == 'настройка модели').sum())),
        (None, None),
        ('Маршрут', 'один: Таллинская (Строгино) ↔ м. Щукинская, ~5,4 км, ~16 остановок'),
        ('Входы', 'ручка от −15 до +15 каждые 0,05 с; скорость передней и задней тележки ~каждые 0,1 с'),
        ('Единицы', 'скорость тележек в данных — в км/ч (не м/с, как в README): делить на 3,6'),
        ('Колёса врут', 'на сколько % скорость тележки больше (+) или меньше (−) реальной по GNSS — износ колёс'),
        ('Проблемы в записи', 'датчик молчал; колёса расходились с GNSS (юз, буксование); сообщения опаздывали'),
    ]
    info.append(['Что в данных', ''])
    for a, b in lines:
        info.append([a, b])
    for row in info.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            c.alignment = Alignment(wrap_text=True, vertical='top')
    style_header(info, [48, 90])
    info.row_dimensions[1].height = 18

    ws = wb.create_sheet('Записи')
    ws.append([h for _, h, _, _ in cols])
    for _, r in cat.iterrows():
        vals = []
        for c, _, _, _ in cols:
            v = r[c]
            if c == 'duration_s':
                v = v / 60
            elif c == 'ref_dist_m' and pd.notna(v):
                v = v / 1000
            vals.append(None if pd.isna(v) or v == '' else (v.item() if hasattr(v, 'item') else v))
        ws.append(vals)
    for row in ws.iter_rows(min_row=2):
        grey = row[1].value in ('не брать', 'отладка')
        for cell, (_, _, fmt, _) in zip(row, cols):
            cell.font = GREY if grey else FONT
            cell.number_format = fmt
        if row[7].value:
            row[7].fill = BAD_FILL
    style_header(ws, [w for _, _, _, w in cols])
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    path = OUT_DIR / 'записи_для_команды.xlsx'
    wb.save(path)
    return path


if __name__ == '__main__':
    arg = sys.argv[1] if len(sys.argv) > 1 else 'catalog'
    print(catalog_xlsx() if arg == 'catalog' else team_xlsx() if arg == 'team' else run_xlsx(arg))
